import marimo

__generated_with = "0.23.9"
app = marimo.App(width="medium")


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    #Horas Extra desde R2-Deltalake
    Fecha: 📅 17 de junio 2026 <br>
    #####Autor: 👨‍💻 Ernesto Palacios <br>
    Objetivo: 🚀 Generar los informes de Horas Extra a partir de los datos procesados de Ordenes de Trabajo. ⛈️
    """)
    return


@app.cell(hide_code=True)
def inicializacion():
    # Importar librerias
    import marimo as mo
    from openpyxl import load_workbook
    from openpyxl.utils.dataframe import dataframe_to_rows
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    from openpyxl.styles import Alignment, numbers
    from openpyxl.utils import get_column_letter

    import shutil
    import os
    import sys
    import time
    from pathlib import Path



    import pandas as pd
    import numpy as np
    import re
    import duckdb

    from natsort import order_by_index, index_natsorted
    from deltalake import DeltaTable, write_deltalake
    from datetime import datetime, time, timedelta
    from datetime import date as toDate
    import calendar
    import locale

    import eerssa.utils
    import eerssa.he_helpers
    import eerssa.excel_styles
    from eerssa.utils import load_r2_credentials
    import eerssa.excel_styles as excel_styles


    import warnings

    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="invalid value encountered in cast")



    ecuador = locale.setlocale(locale.LC_TIME, "es_EC.UTF-8")

    # ── Nombres y Ubicaciones de Archivos Plantilla de Excel ──────────────────────────────────────────────────────────
    t_xls_activ_path = os.path.join('models','plantilla_actividades.xlsx')
    t_xls_consol_path = os.path.join('models','plantilla_consolidado_he.xlsx')
    t_xls_informe_path = os.path.join('models','plantilla_informe_he.xlsx')

    today =  datetime.today().strftime('%Y%m%d')
    new_actividades_file = f"actividades_{today}.xlsx"
    new_base_he_file  = f"base_HE_{today}.xlsx"

    actividades_path = os.path.join('reporte', new_actividades_file)
    base_he_path   = os.path.join('reporte', new_base_he_file )
    return (
        DeltaTable,
        actividades_path,
        calendar,
        dataframe_to_rows,
        datetime,
        duckdb,
        eerssa,
        excel_styles,
        load_r2_credentials,
        load_workbook,
        mo,
        os,
        pd,
        shutil,
        t_xls_activ_path,
        t_xls_consol_path,
        timedelta,
        toDate,
    )


@app.cell(hide_code=True)
def _(mo):
    # Cell: state counter (put this early, near imports)
    get_refresh, set_refresh = mo.state(0)
    return get_refresh, set_refresh


@app.cell
def base_de_datos(DeltaTable, duckdb, load_r2_credentials, mo):
    # Conectar con DELTA LAKE TABLE (Cloudflare R2)
    # DELTA_TABLE_PATH_ON_HOST
    R2_BUCKET = "delta-v30"
    CREDS_PATH = "secrets/r2_credentials.json"
    creds = load_r2_credentials(CREDS_PATH)
    R2_ENDPOINT = f"https://{creds['account_id']}.r2.cloudflarestorage.com"
    table_path =  f"s3://{R2_BUCKET}/delta_v30"

    storage_options = {
        "AWS_ENDPOINT_URL": R2_ENDPOINT,
        "AWS_ACCESS_KEY_ID": creds["access_key"],
        "AWS_SECRET_ACCESS_KEY": creds["secret_key"],
        "AWS_REGION": "auto",
        "AWS_S3_ALLOW_UNSAFE_RENAME": "true",
    }

    # ------  MotherDuck DuckDB  --------- #
    HE_DB_DUCK = f"md:horas_extra?motherduck_token={creds['duckdb_he_token']}"
    con = duckdb.connect(HE_DB_DUCK)

    # ------  Delta Lake  --------- #

    if not DeltaTable.is_deltatable(table_path, storage_options=storage_options):
        mo.stop(
            True,
            mo.callout(
                mo.md(
                    f"""**❌ Error:** No se ha podido conectar la base de datos DELTALAKE en CloudFlare R2 ⛈️ .

    `{table_path}`"""
                ),
                kind="danger",
            ),
        )



    mo.callout(
        mo.md(f"**✅ Conectado** a la tabla Delta Lake en: `{table_path}`"),
        kind="success",
    )
    return con, storage_options, table_path


@app.cell(hide_code=True)
def rango_fechas(calendar, mo, timedelta, toDate):

    # Defaults: first and last day of current month
    _today = toDate.today()
    _first_day = _today.replace(day=1)
    _prev_month_first = (_first_day - timedelta(days=1)).replace(day=1)
    _last_day = _today.replace(day=calendar.monthrange(_today.year, _today.month)[1])

    date_picker = mo.ui.date_range(
        start=_first_day,
        stop=_last_day,
        value=(_first_day, _last_day),
        label="📅 Rango del Reporte",
    )

    mo.md(f"## 1. Seleccione el rango de fechas\n{date_picker}")
    return (date_picker,)


@app.cell(hide_code=True)
def _(mo):
    # Cell: refresh button
    refrescar = mo.ui.run_button(label="🔄 Refrescar datos")
    mo.callout(
        mo.md(f" ❗ Si ha agregado nuevas órdenes de trabajo, por favor refresque los datos:\n\n{refrescar}"),kind="info",
    )
    return (refrescar,)


@app.cell(hide_code=True)
def _(
    DeltaTable,
    date_picker,
    get_refresh,
    mo,
    refrescar,
    storage_options,
    table_path,
):
    # Cell: load data (re-runs on date change OR button click)
    get_refresh()  # dependency — re-runs when state changes
    refrescar.value  # dependency — triggers re-run when clicked

    reporte_inicia, reporte_finaliza = date_picker.value
    _inicio = f"{reporte_inicia}T00:00:00-05:00"
    _fin = f"{reporte_finaliza}T00:00:00-05:00"

    dt = DeltaTable(table_path, storage_options=storage_options)

    df = dt.to_pandas(
        filters=[
            ("Fecha", ">=", _inicio),
            ("Fecha", "<=", _fin),
        ]
    )

    mo.callout(
        mo.md(
            f"**Reporte:** Desde el | _`{reporte_inicia:%A, %d %B %Y}`_ | → hasta el → | _`{reporte_finaliza:%A, %d %B %Y}`_ | 📅  <br> **Version** de la base de datos: #`{dt.version()}`"
        ),
        kind="info",
    )
    return df, dt


@app.cell(hide_code=True)
def _(mo):
    # Cell: Controls
    solo_he = mo.ui.switch(label="Activar para Solo OTs con Horas Extra", value=False)
    btn_crear_excel = mo.ui.run_button(label="📄 Crear archivos Excel")
    return btn_crear_excel, solo_he


@app.cell(hide_code=True)
def _(btn_crear_excel, mo, solo_he):
    # Cell: Display controls
    _label = "**Solo OTs con Horas Extra** 🔥" if solo_he.value else "**Todas las actividades**"

    mo.md(
        f"## 2. Creación de archivo Excel con Actividades\n"
        f"{solo_he} {_label}\n\n"
        f"<br>{btn_crear_excel}"
    )
    return


@app.cell(hide_code=True)
def _(
    actividades_path,
    btn_crear_excel,
    df,
    dt,
    eerssa,
    mo,
    solo_he,
    t_xls_activ_path,
):
    # Cell: Generate Excel (gated)
    mo.stop(
        not btn_crear_excel.value,
        mo.callout(mo.md("⏸️ Presione el botón para crear los archivos Excel."), kind="warn"),
    )

    if solo_he.value:
        # Keep only OTs where at least one row has HorasExtra == 'Si'
        _ots_con_he = df.loc[df['HorasExtra'] == 'Si', 'id_ot'].unique()
        df_export = df[df['id_ot'].isin(_ots_con_he)]
    else:
        df_export = df.copy()

    n_rows = eerssa.he_helpers.exportar_actividades_excel(
        df=df_export,
        plantilla_path=t_xls_activ_path,
        output_path=actividades_path,
    )

    _mode = "Solo Horas Extra 🔥" if solo_he.value else "Todas las actividades"
    mo.callout(
        mo.md(
            f"✅ **Excel generado** — `{n_rows}` filas ({_mode})"
            f"<br>Version DeltaLake: `{dt.version()}`"
            f"<br>Archivo: `{actividades_path}`"
        ),
        kind="success",
    )
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## 3. Recuperar actividades desde Excel
    Una vez se han editado y limpiado los datos de actividades en el archivo Excel se procede a cargar los datos y posterior actualizar la base de datos
    """)
    return


@app.cell(hide_code=True)
def _(mo):
    btn_leer_excel = mo.ui.run_button(label="👓 Cargar cambios desde Excel Actividades") 
    btn_leer_excel
    return (btn_leer_excel,)


@app.cell(hide_code=True)
def _(actividades_path, btn_leer_excel, df, eerssa, mo, pd):
    # 1. Load the modified Excel file: ACTIVIDADES_2026XXXXXX

    btn_leer_excel.value

    excel_file_path = actividades_path
    # Aqui puedo escoger recargar desde algun archivo Excel para pruebas,
    # por defecto es el generado en el paso anterior

    modified_df = pd.read_excel(excel_file_path, sheet_name = "ACTIVIDADES")
    modified_df = modified_df.dropna(subset=['Item'])

    # 2. Se vuelva a colocar el String de TimeZone en la Fecha
    modified_df['Fecha'] = modified_df['Fecha'].apply(lambda x: eerssa.he_helpers.ColocarTimezone( x ))

    # 3. Convertir de String a TimeObject y se vuelve a calcular la duración en minutos
    modified_df['Ini'] = pd.to_datetime(modified_df['InicioEvento'], errors='coerce')
    modified_df['Fin'] = pd.to_datetime(modified_df['FinEvento'], errors='coerce')

    modified_df['Duracion'] = eerssa.he_helpers.calcular_minutos_transcurridos(
        modified_df['Ini'],
        modified_df['Fin']
    )

    #4. Convert everything to datetime objects first, then format them all as uniform strings
    modified_df['InicioEvento'] = pd.to_datetime(modified_df['InicioEvento'], errors='coerce').dt.strftime('%Y-%m-%d %H:%M:%S')
    modified_df['FinEvento'] = pd.to_datetime(modified_df['FinEvento'], errors='coerce').dt.strftime('%Y-%m-%d %H:%M:%S')

    # 5. Ensure Schema Consistency
    # Excel often introduces new columns (like empty comments) or reorders them.
    # We force the modified_df to have the same columns as the original df.

    modified_df = modified_df.drop(columns=['Ini', 'Fin'])
    modified_df = modified_df[list(df.columns)]

    # Visualize the changes
    mo.md(f"### Vista previa: `{len(modified_df)}` filas a guardar")
    mo.ui.table(modified_df)
    return (modified_df,)


@app.cell(hide_code=True)
def _(mo, modified_df):
    btn_guardar = mo.ui.run_button(label="💾 Guardar en Delta Lake", kind="danger")
    mo.md(
        f"## 4. Actualizar Delta-Lake\n"
        f"`{len(modified_df)}` filas pendientes de guardar\n\n"
        f"{btn_guardar}"
    )
    return (btn_guardar,)


@app.cell(hide_code=True)
def _(btn_guardar, dt, get_refresh, mo, modified_df, set_refresh):
    # Con el archivo modificado en Excel, se actualizan las filas en DeltaLake
    mo.stop(not btn_guardar.value, mo.callout(mo.md("⏸️ Revise los datos antes de guardar."), kind="warn"))
    # ... your merge code ...
    try:
        # --- Start of new logic ---
        # 1. Get a list of all unique 'id_ot' values from the source DataFrame.
        ids_to_update = modified_df['id_ot'].unique()

        # 2. Format the list into a SQL-compatible string like "(101, 102, 103)".
        # This is crucial for the IN clause to work correctly.
        ids_predicate_string = ", ".join(map(str, ids_to_update))

        # 3. Define the delete predicate to scope deletions to only the OTs being updated.
        delete_predicate = f"target.id_ot IN ({ids_predicate_string})"
        # --- End of new logic ---

        # The unique key for matching rows remains the same.
        unique_key_predicate = "target.id_ot = source.id_ot AND target.Item = source.Item"

        (dt.merge(
                source=modified_df,
                predicate=unique_key_predicate,
                source_alias="source",
                target_alias="target"
            )
            .when_matched_update_all()  # Rule 1: If a row exists, update it.
            .when_not_matched_insert_all()  # Rule 2: If it's a new row, insert it.
            .when_not_matched_by_source_delete(  # Rule 3: If an old row is now gone...
                predicate=delete_predicate  # ...delete it, but ONLY if it belongs to an OT we are modifying.
            )
            .execute()
        )
        # ✅ Bump the refresh counter → triggers df reload
        set_refresh(get_refresh() + 1)

        saved = "✅ **Successfully saved changes for all modified OTs to Delta Lake!**"
    except Exception as e:
        saved = f"❌ **Error saving to Delta Lake:** {e}"


    mo.callout(mo.md(f"{saved}<br>**VERSION Actual:** `{dt.version()}`"), kind="success")
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    --------------------
    """)
    return


@app.cell
def _(df, eerssa, t_xls_consol_path):
    reglas = eerssa.he_helpers.cargar_reglas_tipo(t_xls_consol_path)
    consolidado = eerssa.he_helpers.consolidar_horas_extra(df, reglas)
    return (consolidado,)


@app.cell
def paso_2_enriquecer(con, consolidado, date_picker, eerssa, mo):

    _inicio, _fin = date_picker.value

    consolidado_enriquecido, n_coincidencias = eerssa.he_helpers.enriquecer_desde_duckdb(
        con, consolidado, _inicio, _fin
    )

    _total = len(consolidado_enriquecido)
    _nuevos = _total - n_coincidencias

    mo.callout(
        mo.md(
            f"**Paso 2 — Enriquecimiento desde DuckDB**<br>"
            f"Total: `{_total}` filas · "
            f"Coincidencias: `{n_coincidencias}` (ediciones preservadas) · "
            f"Nuevos: `{_nuevos}` (auto-generados)"
        ),
        kind="info" if n_coincidencias > 0 else "neutral",
    )
    return (consolidado_enriquecido,)


@app.cell
def _(consolidado_enriquecido):
    consolidado_enriquecido
    return


@app.cell
def paso_3_controles(mo):
    btn_exportar_consol = mo.ui.run_button(label="📤 Exportar Consolidado a Excel")
    mo.md(
        f"### Paso 3 — Exportar Consolidado\n"
        f"Exporte el consolidado enriquecido a Excel para revisar **Evento** y **Cuenta**.\n\n"
        f"{btn_exportar_consol}"
    )
    return (btn_exportar_consol,)


@app.cell
def paso_3_exportar(
    btn_exportar_consol,
    consolidado_enriquecido,
    dataframe_to_rows,
    datetime,
    excel_styles,
    load_workbook,
    mo,
    os,
    shutil,
    t_xls_consol_path,
):
    mo.stop(
        not btn_exportar_consol.value,
        mo.callout(mo.md("⏸️ Presione el botón para exportar."), kind="warn"),
    )

    _today = datetime.today().strftime('%Y%m%d')
    _path = os.path.join('reporte', f"consolidado_he_{_today}.xlsx")

    shutil.copyfile(t_xls_consol_path, _path)

    _wb = load_workbook(_path)
    _ws = _wb.worksheets[1]

    # Clear existing data before writing
    for _row in _ws.iter_rows(min_row=2, max_row=_ws.max_row, min_col=1, max_col=_ws.max_column):
        for _cell in _row:
            _cell.value = None

    # Duracion en Formato Excel
    _consolidado_excel = consolidado_enriquecido.copy()
    _consolidado_excel['Duracion'] = _consolidado_excel['Duracion'] / (60*24)

    for r_idx, _row in enumerate(dataframe_to_rows(_consolidado_excel, index=False, header=False), 2):
        for c_idx, value in enumerate(_row, 1):
            _ws.cell(row=r_idx, column=c_idx, value=value)

    excel_styles.apply_conditional_formatting(_ws)

    _wb.save(_path)

    mo.callout(
        mo.md(f"✅ **Excel exportado:** `{_path}`<br>Edite **Evento** y **Cuenta**. Formato: `REDES:50, MEDIDORES:50`"),
        kind="success",
    )
    return


@app.cell
def paso_3b_controles(mo):
    btn_leer_consol = mo.ui.run_button(label="👓 Cargar Consolidado editado")
    mo.md(
        f"### Paso 3b — Cargar Excel editado\n"
        f"Una vez terminada la revisión en Excel:\n\n"
        f"{btn_leer_consol}"
    )
    return (btn_leer_consol,)


@app.cell
def paso_3b_leer(btn_leer_consol, datetime, eerssa, load_workbook, mo, os, pd):
    mo.stop(
        not btn_leer_consol.value,
        mo.callout(mo.md("⏸️ Presione el botón para cargar el Excel editado."), kind="warn"),
    )

    _today = datetime.today().strftime('%Y%m%d')
    _path = os.path.join('reporte', f"consolidado_he_{_today}.xlsx")

    _wb = load_workbook(_path, data_only=False)
    _ws = _wb["result"]
    data = _ws.values
    cols = next(data)
    consolidado_editado = pd.DataFrame(data, columns=cols)
    consolidado_editado = consolidado_editado.dropna(subset=['Cuadrilla'])

    consolidado_editado['Cuenta'] = consolidado_editado['Cuenta'].apply(
        eerssa.he_helpers.cuenta_str_to_map
    )
    consolidado_editado['id_ot'] = consolidado_editado['id_ot'].apply(lambda x: int(x))
    consolidado_editado['Fecha'] = pd.to_datetime(consolidado_editado['Fecha']).dt.date
    consolidado_editado['InicioEvento'] = pd.to_datetime(
        consolidado_editado['InicioEvento'], format='%H:%M:%S'
    ).dt.time
    consolidado_editado['FinEvento'] = pd.to_datetime(
        consolidado_editado['FinEvento'], format='%H:%M:%S'
    ).dt.time
    consolidado_editado['Duracion'] = (
        pd.to_datetime(consolidado_editado['FinEvento'].astype(str))
        - pd.to_datetime(consolidado_editado['InicioEvento'].astype(str))
    ).dt.total_seconds() // 60

    mo.callout(
        mo.md(f"✅ **Excel cargado:** `{len(consolidado_editado)}` filas"),
        kind="success",
    )
    mo.ui.table(consolidado_editado)
    return (consolidado_editado,)


@app.cell
def _(consolidado_editado, mo):
    btn_sync_duck = mo.ui.run_button(label="🦆 Sincronizar DuckDB", kind="danger")
    mo.md(
        f"## 5. Sincronizar con DuckDB (MotherDuck)\n"
        f"`{len(consolidado_editado)}` filas pendientes de sincronizar\n\n"
        f"{btn_sync_duck}"
    )
    return (btn_sync_duck,)


@app.cell
def _(btn_sync_duck, con, consolidado_editado, date_picker, eerssa, mo):
    mo.stop(
        not btn_sync_duck.value,
        mo.callout(mo.md("⏸️ Presione 🦆 para sincronizar con DuckDB."), kind="warn"),
    )

    _inicio, _fin = date_picker.value

    n_act, n_part = eerssa.he_helpers.sync_deltalake_to_duckdb(
        con, consolidado_editado, _inicio, _fin
    )

    mo.callout(
        mo.md(
            f"✅ **DuckDB sincronizado**"
            f"<br>Actividades: `{n_act}` — Participaciones: `{n_part}`"
        ),
        kind="success",
    )
    return


@app.cell
def paso_5b_gdrive_validacion(con, date_picker, mo):
    import eerssa.organizar as gdrive

    _inicio, _fin = date_picker.value

    cuadrilla_df = gdrive.get_gsheet_df()

    _initials = con.execute("""
        SELECT DISTINCT p.Responsable
        FROM participaciones p
        JOIN actividades a ON a.id_actividad = p.id_actividad
        WHERE a.Fecha BETWEEN $1 AND $2
    """, [str(_inicio), str(_fin)]).df()

    _initials_list = _initials['Responsable'].tolist() if not _initials.empty else []
    _missing = gdrive.validar_iniciales(_initials_list, cuadrilla_df)

    if _missing:
        mo.callout(
            mo.md(
                f"⚠️ **Iniciales no encontradas en Google Drive:** {', '.join(_missing)}<br>"
                f"Actualice el documento de Google Drive (`DB_calificar_ot`) para agregar estas personas."
            ),
            kind="warn",
        )
    return cuadrilla_df, gdrive


@app.cell
def paso_5b_cuadrilla_dropdown(cuadrilla_df, gdrive, mo):
    _cuadrillas = gdrive.get_lista_cuadrillas_ordenadas(cuadrilla_df)

    _opciones = {corto: corto for _ot, corto in _cuadrillas}
    cuadrilla_select = mo.ui.dropdown(
        options=_opciones,
        label="Cuadrilla",
        value=None,
    )

    mo.md(
        f"## 5b. Ajustes individuales de tiempo\n"
        f"Seleccione la Cuadrilla:\n\n"
        f"{cuadrilla_select}"
    )
    return (cuadrilla_select,)


@app.cell
def paso_5b_persona_dropdown(cuadrilla_df, cuadrilla_select, gdrive, mo):
    _selected_c = cuadrilla_select.value

    if _selected_c is None:
        mo.callout(mo.md("⏸️ Seleccione una Cuadrilla para ver el personal."), kind="warn")
        persona_select = mo.ui.dropdown(options={}, label="Persona")
    else:
        _nombres = gdrive.get_personal_cuadrilla(_selected_c, cuadrilla_df)
        persona_select = mo.ui.dropdown(
            options={n: n for n in _nombres},
            label="Persona",
        )

    mo.md(f"{persona_select}")
    return (persona_select,)


@app.cell
def paso_5b_control(mo):
    btn_revisar = mo.ui.run_button(label="🔍 Revisar")
    mo.md(
        f"Una vez seleccionada la persona:\n\n"
        f"{btn_revisar}"
    )
    return (btn_revisar,)


@app.cell
def paso_5b_query(
    btn_revisar,
    con,
    cuadrilla_df,
    date_picker,
    mo,
    persona_select,
):
    mo.stop(
        not btn_revisar.value,
        mo.callout(mo.md("⏸️ Presione **Revisar** para consultar los tiempos de la persona."), kind="warn"),
    )

    _inicio, _fin = date_picker.value
    _persona = persona_select.value

    # 1. Use mo.stop instead of return
    mo.stop(
        _persona is None,
        mo.callout(mo.md("⚠️ Seleccione una persona primero."), kind="warn")
    )

    _match = cuadrilla_df[cuadrilla_df['NOMBRE'] == _persona]

    # 2. Use mo.stop instead of return
    mo.stop(
        _match.empty,
        mo.callout(mo.md(f"⚠️ No se encontró `{_persona}` en Google Drive."), kind="warn")
    )

    iniciales_persona = _match['INICIALES'].iloc[0]

    df_tiempos = con.execute("""
        SELECT a.id_actividad,
               p.Responsable,
               a.Dia,
               a.Fecha,
               COALESCE(p.InicioEvento, a.InicioEvento) AS InicioEvento,
               COALESCE(p.FinEvento, a.FinEvento) AS FinEvento,
               a.Evento,
               a.Duracion,
               a.Cuenta,
               a.id_ot,
               a.Items,
               a.Tipo,
               a.Cuadrilla,
               p.tiempo_ajustado
        FROM actividades a
        JOIN participaciones p ON a.id_actividad = p.id_actividad
        WHERE a.Fecha BETWEEN $1 AND $2
          AND p.Responsable = $3
        ORDER BY a.Fecha, a.InicioEvento
    """, [str(_inicio), str(_fin), iniciales_persona]).df()

    mo.stop(
        df_tiempos.empty,
        mo.callout(
            mo.md(f"ℹ️ No hay participaciones para `{_persona}` ({iniciales_persona}) en el rango seleccionado."),
            kind="neutral",
        )
    )

    df_tiempos['Fecha'] = df_tiempos['Fecha'].apply(lambda x: str(x)[:10])
    df_tiempos['InicioEvento'] = df_tiempos['InicioEvento'].astype(str)
    df_tiempos['FinEvento'] = df_tiempos['FinEvento'].astype(str)

    _cols_mostrar = ['Responsable', 'Dia', 'Fecha', 'InicioEvento', 'FinEvento', 'Evento']
    editable_df = mo.ui.data_editor(df_tiempos[_cols_mostrar])

    mo.vstack([
        mo.callout(
            mo.md(f"**{iniciales_persona}** — {_persona} — {len(df_tiempos)} participaciones"),
            kind="info",
        ),
        mo.md("✏️ **Haga doble clic en una celda para editar los tiempos:**"),
        editable_df,
    ])
    return df_tiempos, editable_df


@app.cell
def paso_5b_guardar_button(mo):
    btn_guardar_tiempos = mo.ui.run_button(label="💾 Guardar ajustes", kind="danger")
    mo.md(
        f"Una vez editados los tiempos:\n\n"
        f"{btn_guardar_tiempos}"
    )
    return (btn_guardar_tiempos,)


@app.cell
def paso_5b_guardar_exec(
    btn_guardar_tiempos,
    con,
    df_tiempos,
    editable_df,
    mo,
):
    mo.stop(
        not btn_guardar_tiempos.value,
        mo.callout(mo.md("⏸️ Presione **Guardar ajustes** para guardar los cambios en la base de datos."), kind="warn"),
    )

    _edited = editable_df.value

    n_updated = 0
    for i, _row in _edited.iterrows():
        _original = df_tiempos.loc[i]
        _nuevo_inicio = _row['InicioEvento']
        _nuevo_fin = _row['FinEvento']

        if _nuevo_inicio != _original['InicioEvento'] or _nuevo_fin != _original['FinEvento']:
            _id_act = int(_original['id_actividad'])
            con.execute("""
                UPDATE participaciones SET
                    InicioEvento = $1::TIME,
                    FinEvento    = $2::TIME,
                    tiempo_ajustado = TRUE
                WHERE id_actividad = $3 AND Responsable = $4
            """, [_nuevo_inicio, _nuevo_fin, _id_act, _original['Responsable']])
            n_updated += 1

    mo.callout(
        mo.md(f"✅ **Ajustes guardados:** `{n_updated}` filas actualizadas en DuckDB."),
        kind="success",
    )
    return


@app.cell
def paso_6_controles(mo):
    btn_generar_informe = mo.ui.run_button(label="📄 Generar Informe Final", kind="danger")
    mo.md(
        f"## 6. Generar Informe Final\n"
        f"Genera el reporte final con una hoja por persona.\n\n"
        f"{btn_generar_informe}"
    )
    return (btn_generar_informe,)


@app.cell
def paso_6_generar(btn_generar_informe, con, date_picker, eerssa, mo):
    mo.stop(
        not btn_generar_informe.value,
        mo.callout(mo.md("⏸️ Presione el botón para generar el informe."), kind="warn"),
    )

    _inicio, _fin = date_picker.value
    _template = "models/plantilla_informe_he.xlsx"
    _output = "reporte/00_Informe_HE_todos.xlsx"

    personas = eerssa.he_helpers.generar_informe_he(
        con, _inicio, _fin, _template, _output
    )

    mo.callout(
        mo.md(
            f"✅ **Informe generado:** `{_output}`<br>"
            f"Personas: `{len(personas)}` — {', '.join(personas)}"
        ),
        kind="success",
    )
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    📎
    """)
    return


if __name__ == "__main__":
    app.run()
