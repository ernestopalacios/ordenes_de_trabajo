from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule


# ============================================================
# STYLE DEFINITIONS
# ============================================================

style_fin_evento = DifferentialStyle(
    font=Font(color="9C0006"),
    fill=PatternFill(bgColor="FFC7CE"),
)

style_sabado = DifferentialStyle(
    font=Font(color="000000"),
    fill=PatternFill(bgColor="FFEB9C"),
)

style_domingo = DifferentialStyle(
    font=Font(color="000000"),
    fill=PatternFill(bgColor="EEE31D"),
)

style_green = DifferentialStyle(
    font=Font(color="006100"),
    fill=PatternFill(bgColor="C6EFCE"),
)

style_red_error = DifferentialStyle(
    font=Font(color="FFFFFF"),
    fill=PatternFill(bgColor="C00000"),
)

style_red_text = DifferentialStyle(
    font=Font(color="000000"),
    fill=PatternFill(bgColor="FF9D9D"),
)

style_bold = DifferentialStyle(
    font=Font(bold=True),
)

style_blue = DifferentialStyle(
    font=Font(color="FFFFFF"),
    fill=PatternFill(bgColor="4472C4"),
)

style_mad = DifferentialStyle(
    font=Font(color="000000"),
    fill=PatternFill(bgColor="97B4E8"),
)

style_festivo = DifferentialStyle(
    font=Font(color="FFFFFF"),
    fill=PatternFill(bgColor="7030A0"),
)

style_purple = DifferentialStyle(
    font=Font(color="FFFFFF"),
    fill=PatternFill(bgColor="FF8D8DB3"),
)

style_gray = DifferentialStyle(
    font=Font(color="FFFFFF"),
    fill=PatternFill(bgColor="808080"),
)

style_bottom_border = DifferentialStyle(
    border=Border(bottom=Side(style="thick", color="000000")),
)


# ============================================================
# CONDITIONAL FORMATTING RULES
# ============================================================

rules_informe_he = [
    {
        "formula": '=OR(WEEKDAY($B6, 2) > 5, ISNUMBER(SEARCH("FESTIVO", $J6)))',
        "style": style_purple,
        "range": "A6:J36",
    },
    {
        "formula": '=ISNUMBER(SEARCH("FESTIVO", $J6))',
        "style": style_bold,
        "range": "A6:I36",
    },
]

rules = [
    {
        "formula": '=COUNTIF(Revisar_primero!$B$2:$B$14,$D1)>0',
        "style": style_festivo,
        "range": "A1:N5000",
    },
    {
        "formula": '=AND(OR($A1="Zamora Z1 (Cuadrilla. Nro. 6)",$A1="Zamora Z1 (Cuadrilla. AP Nro. 4)",$A1="Zamora (Agencia)",$A1="Jefatura Zonal Zamora"),COUNTIF(Revisar_primero!$B$15:$B$18,$D1)>0)',
        "style": style_festivo,
        "range": "A1:N5000",
    },
    {
        "formula": '=AND($A1="Yacuambi Z1 (Cuadrilla. Nro. 8)",COUNTIF(Revisar_primero!$B$19:$B$19,$D1)>0)',
        "style": style_festivo,
        "range": "A1:N5000",
    },
    {
        "formula": '=AND(OR($A1="Yantzaza Z1 (Cuadrilla. Nro. 5)",$A1="Lineas Energizadas (Cuadrilla Nro.6)",$A1="Yantzaza (Agencia)"),COUNTIF(Revisar_primero!$B$20:$B$22,$D1)>0)',
        "style": style_festivo,
        "range": "A1:N5000",
    },
    {
        "formula": '=AND($A1="Paquisha Z1 (Cuadrilla. Nro. 10)",COUNTIF(Revisar_primero!$B$23:$B$23,$D1)>0)',
        "style": style_festivo,
        "range": "A1:N5000",
    },
    {
        "formula": '=AND($A1="Guayzimi Z1 (Cuadrilla. Nro. 7)",COUNTIF(Revisar_primero!$B$24:$B$24,$D1)>0)',
        "style": style_festivo,
        "range": "A1:N5000",
    },
    {
        "formula": '=AND(OR($A1="El Pangui Z1 (Cuadrilla. Nro. 4)",$A1="El Pangui (Agencia)"),COUNTIF(Revisar_primero!$B$25:$B$26,$D1)>0)',
        "style": style_festivo,
        "range": "A1:N5000",
    },
    {
        "formula": '=AND(OR($A1="Gualaquiza Z1 (Cuadrilla. Nro. 3)",$A1="Gualaquiza (Agencia)"),COUNTIF(Revisar_primero!$B$27:$B$28,$D1)>0)',
        "style": style_festivo,
        "range": "A1:N5000",
    },
    {
        "formula": '=$C1="sábado"',
        "style": style_sabado,
        "range": "A1:N5000",
    },
    {
        "formula": '=$C1="domingo"',
        "style": style_domingo,
        "range": "A1:N5000",
    },
    {
        "formula": '=AND($A1="Zamora Z1 (Cuadrilla. AP Nro. 4)",COUNTIF(Revisar_primero!$D$2:$D$114,$D1)>0)',
        "style": style_blue,
        "range": "A1:N5000",
    },
    {
        "formula": "=$A1<>$A2",
        "style": style_bottom_border,
        "range": "A1:N5000",
    },
    {
        "formula": '=AND($E1<>"",(HOUR($E1)<6),((HOUR($F1)+MINUTE($F1)/100)>6))',
        "style": style_red_error,
        "range": "A1:N5000",
    },
    {
        "formula": "=$F1<$E1",
        "style": style_red_error,
        "range": "A1:N5000",
    },
    {
        "formula": '=AND($E1<>"",HOUR($E1)<6)',
        "style": style_mad,
        "range": "A1:N5000",
    },
    {
        "formula": "=AND(HOUR($E1)<8,(HOUR($F1)+MINUTE($F1)/100)>8)",
        "style": style_red_text,
        "range": "F1:F5000",
    },
    {
        "formula": "=AND(HOUR($E1)<17,HOUR($F1)>17)",
        "style": style_red_text,
        "range": "E1:E5000",
    },
]


# ============================================================
# APPLY FUNCTIONS
# ============================================================

def apply_conditional_formatting(ws):
    """Apply all conditional formatting rules to a worksheet."""
    for r in rules:
        rule = Rule(
            type="expression",
            formula=[r["formula"]],
            dxf=r["style"],
            stopIfTrue=False,
        )
        ws.conditional_formatting.add(r["range"], rule)


def apply_he_informe(ws):
    """Apply HE inform rules to a worksheet."""
    for r in rules_informe_he:
        rule = Rule(
            type="expression",
            formula=[r["formula"]],
            dxf=r["style"],
            stopIfTrue=False,
        )
        ws.conditional_formatting.add(r["range"], rule)